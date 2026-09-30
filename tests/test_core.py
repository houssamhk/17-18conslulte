import os
import sqlite3
import tempfile
import unittest
from datetime import datetime, timedelta
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import bcrypt
from cryptography.fernet import Fernet

from engine.anonymizer import Anonymizer
from engine.compliance_engine import AlgComplianceEngine
from engine.document_parser import DocumentParser
from engine.license_manager import LicenseManager
from engine.nlp_detector import NLPDetector
from engine.regex_detector import RegexDetector
from engine.scheduler import TaskScheduler
from storage.key_manager import KeyManager
from storage.secure_db import SecureDatabase
from storage.secure_vault import SecureVault


class CoreFeatureTests(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory()
        self.key = Fernet.generate_key()
        self.key_patch = patch.object(KeyManager, "ensure_key", return_value=self.key)
        self.key_patch.start()
        self.addCleanup(self.key_patch.stop)
        # Unit tests use an isolated in-memory stand-in; Windows Credential
        # Manager is unavailable in some CI/service sessions (CredWrite 1312).
        self.audit_heads = {}
        def store_audit_head(value, account=KeyManager.AUDIT_HEAD_ACCOUNT):
            self.audit_heads[account] = value
            return True

        self.audit_store_patch = patch.object(
            KeyManager,
            "store_audit_head",
            side_effect=store_audit_head,
        )
        self.audit_retrieve_patch = patch.object(
            KeyManager,
            "retrieve_audit_head",
            side_effect=lambda account=KeyManager.AUDIT_HEAD_ACCOUNT: self.audit_heads.get(account),
        )
        self.audit_store_patch.start()
        self.audit_retrieve_patch.start()
        self.addCleanup(self.audit_store_patch.stop)
        self.addCleanup(self.audit_retrieve_patch.stop)
        self.addCleanup(self.tempdir.cleanup)

    def database(self, name="test.db"):
        return SecureDatabase(os.path.join(self.tempdir.name, name))

    def test_regex_detection_and_redaction(self):
        source = "اتصل على 0551234567 أو راسل test@example.dz"
        entities = RegexDetector().detect(source)
        self.assertEqual({entity.entity_type for entity in entities}, {"PHONE", "EMAIL"})
        output, highlighted = Anonymizer().anonymize(source, entities, "full_mask")
        self.assertNotIn("0551234567", output)
        self.assertNotIn("test@example.dz", output)
        self.assertIn("<span", highlighted)

    def test_detector_covers_supported_identifier_patterns(self):
        text = """123456789012345678 0551234567 1234567890 12 12345678901234567890
        123456789012345 DZ123456789012345678901234 person@example.dz A12345678"""
        entities = RegexDetector().detect(text)
        self.assertEqual(
            {entity.entity_type for entity in entities},
            {"NIN", "PHONE", "CCP", "RIB", "NIF", "IBAN", "EMAIL", "PASSPORT"},
        )

    def test_optional_nlp_never_downloads_missing_weights(self):
        detector = NLPDetector(os.path.join(self.tempdir.name, "missing-model"))
        with self.assertLogs(level="ERROR"):
            self.assertFalse(detector.load_model())
        self.assertFalse(detector.is_available())

    def test_license_signing_uses_rsa_private_key_and_rejects_legacy_hmac(self):
        from cryptography.hazmat.primitives import serialization
        from cryptography.hazmat.primitives.asymmetric import rsa
        import jwt
        from engine.license_signer import generate_license

        private_key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
        private_path = os.path.join(self.tempdir.name, "issuer.pem")
        with open(private_path, "wb") as key_file:
            key_file.write(private_key.private_bytes(
                serialization.Encoding.PEM,
                serialization.PrivateFormat.PKCS8,
                serialization.NoEncryption(),
            ))
        public_key = private_key.public_key().public_bytes(
            serialization.Encoding.PEM,
            serialization.PublicFormat.SubjectPublicKeyInfo,
        ).decode("ascii")

        with patch("engine.license_signer._private_key_path", return_value=__import__("pathlib").Path(private_path)), \
             patch.object(LicenseManager, "PUBLIC_KEY", public_key):
            token = generate_license(LicenseManager.get_hardware_id(), 30)
            self.assertTrue(LicenseManager.validate_license(token)[0])
            wrong_hwid = generate_license("0000000000000000", 30)
            self.assertFalse(LicenseManager.validate_license(wrong_hwid)[0])

        old_token = jwt.encode(
            {"hwid": LicenseManager.get_hardware_id(), "product": "Alg-PII Engine Enterprise"},
            "legacy-test-signing-key-that-is-at-least-thirty-two-bytes-long",
            algorithm="HS256",
        )
        self.assertFalse(LicenseManager.validate_license(old_token)[0])

    def test_license_issuer_database_encrypts_legacy_client_records(self):
        import sqlite3
        from keygen_app import LicenseDB

        path = os.path.join(self.tempdir.name, "issuer.db")
        plain = sqlite3.connect(path)
        with plain:
            plain.execute("""CREATE TABLE clients (
                id INTEGER PRIMARY KEY AUTOINCREMENT, company_name TEXT NOT NULL,
                contact_email TEXT, hwid TEXT NOT NULL UNIQUE, license_key TEXT,
                issue_date DATETIME, expiry_date DATETIME, status TEXT DEFAULT 'ACTIVE')""")
            plain.execute(
                "INSERT INTO clients (company_name, contact_email, hwid) VALUES (?, ?, ?)",
                ("Example Co", "privacy@example.dz", "0011223344556677"),
            )
        plain.close()
        with open(path, "rb") as source:
            original_bytes = source.read()

        credentials = {}
        with patch("keygen_app.keyring.get_password", side_effect=lambda service, account: credentials.get((service, account))), \
             patch("keygen_app.keyring.set_password", side_effect=lambda service, account, value: credentials.__setitem__((service, account), value)):
            db = LicenseDB(path)
            self.assertEqual(db.get_clients()[0][1:3], ("Example Co", "0011223344556677"))
            db.close()

        with open(path, "rb") as encrypted_file:
            self.assertNotEqual(encrypted_file.read(16), b"SQLite format 3\x00")
        backup_path = path + ".legacy-backup.fernet"
        with open(backup_path, "rb") as backup:
            backup_key = credentials[(LicenseDB.SERVICE_NAME, LicenseDB.KEY_ACCOUNT)].encode("ascii")
            self.assertEqual(Fernet(backup_key).decrypt(backup.read()), original_bytes)

    def test_pseudonyms_are_stable_and_not_real_identifiers(self):
        detector = RegexDetector()
        entities = detector.detect("0551234567 ثم 0551234567")
        anonymizer = Anonymizer()
        masked, _ = anonymizer.anonymize("0551234567 ثم 0551234567", entities, "pseudonymize")
        values = [part for part in masked.split() if part.startswith("[PHONE-")]
        self.assertEqual(values[0], values[1])
        self.assertNotIn("0550000000", masked)

    def test_compliance_report_has_review_warning(self):
        engine = AlgComplianceEngine(use_nlp=False)
        report = engine.generate_report(engine.scan("0551234567"))
        self.assertTrue(any("qualified Algerian counsel" in item for item in report.recommendations))
        self.assertEqual(report.risk_level, "LOW")

    def test_encrypted_database_auth_history_and_retention(self):
        db_path = os.path.join(self.tempdir.name, "secure.db")
        db = SecureDatabase(db_path)
        self.addCleanup(db.close)
        self.assertTrue(db.create_user("rootuser", "a-secure-password", "admin"))
        self.assertEqual(db.authenticate_user("rootuser", "a-secure-password")[0], True)
        self.assertEqual(db.authenticate_user("rootuser", "incorrect")[0], False)
        db.log_scan("fixture.txt", 1, "LOW", "full_mask", {"PHONE": 1})
        self.assertEqual(db.get_scan_history()[0][2], "fixture.txt")
        self.assertFalse(db.create_user("weak", "short", "admin"))

        plain = sqlite3.connect(db_path)
        try:
            with self.assertRaises(sqlite3.DatabaseError):
                plain.execute("SELECT * FROM users").fetchall()
        finally:
            plain.close()

        db.save_retention_policy("scan_history", 1)
        db.conn.execute("UPDATE scan_history SET timestamp=datetime('now', '-3 days')")
        db.conn.commit()
        db.run_retention_purge()
        self.assertEqual(db.get_scan_history(), [])

    def test_plain_database_migrates_with_encrypted_backup(self):
        path = os.path.join(self.tempdir.name, "legacy.db")
        source_text = "legacy-data"
        legacy_key = Fernet(self.key)
        with sqlite3.connect(path) as legacy:
            legacy.execute("""CREATE TABLE scan_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT, timestamp DATETIME DEFAULT CURRENT_TIMESTAMP,
                document_name TEXT, total_entities INTEGER, risk_level TEXT, strategy_used TEXT,
                summary_json TEXT, sensitivity_label TEXT DEFAULT 'UNCLASSIFIED', ocr_used BOOLEAN DEFAULT 0,
                duplicates_found INTEGER DEFAULT 0, label_override_by TEXT, label_override_reason TEXT,
                department TEXT)""")
            legacy.execute("CREATE TABLE users (id INTEGER PRIMARY KEY, username TEXT, password_hash TEXT, role TEXT)")
            legacy.execute("INSERT INTO users VALUES (1, 'admin', ?, 'admin')", (
                bcrypt.hashpw(b"admin", bcrypt.gensalt()).decode(),
            ))
            legacy.execute("INSERT INTO scan_history (document_name, total_entities, risk_level, strategy_used, summary_json) VALUES (?, 1, 'LOW', 'full_mask', ?)", (
                legacy_key.encrypt(source_text.encode()).decode(),
                legacy_key.encrypt(b'{"PHONE": 1}').decode(),
            ))
        legacy.close()
        with open(path, "rb") as original:
            original_bytes = original.read()

        db = SecureDatabase(path)
        self.addCleanup(db.close)
        self.assertEqual(db.get_scan_history()[0][2], source_text)
        self.assertFalse(db.has_users())
        backup_path = path + ".legacy-backup.fernet"
        self.assertTrue(os.path.isfile(backup_path))
        with open(backup_path, "rb") as backup:
            self.assertEqual(legacy_key.decrypt(backup.read()), original_bytes)

    def test_legacy_schema_adds_dashboard_and_audit_columns(self):
        path = os.path.join(self.tempdir.name, "old-schema.db")
        legacy = sqlite3.connect(path)
        with legacy:
            legacy.execute("CREATE TABLE scan_history (id INTEGER PRIMARY KEY, timestamp DATETIME, document_name TEXT, total_entities INTEGER, risk_level TEXT, strategy_used TEXT, summary_json TEXT)")
            legacy.execute("CREATE TABLE audit_log (id INTEGER PRIMARY KEY, timestamp DATETIME, action TEXT, details TEXT)")
            legacy.execute("INSERT INTO scan_history VALUES (1, CURRENT_TIMESTAMP, 'old.txt', 1, 'LOW', 'full_mask', '{}')")
        legacy.close()
        db = SecureDatabase(path)
        self.addCleanup(db.close)
        self.assertEqual(db.get_scan_history()[0][7], "UNCLASSIFIED")
        db.log_audit("adminuser", "TEST", "legacy schema")
        self.assertEqual(db.get_audit_log()[0][1:], ("adminuser", "TEST", "legacy schema"))

    def test_legacy_totp_foreign_key_does_not_block_bootstrap_removal(self):
        path = os.path.join(self.tempdir.name, "legacy-totp.db")
        legacy = sqlite3.connect(path)
        with legacy:
            legacy.execute("CREATE TABLE users (id INTEGER PRIMARY KEY, username TEXT, password_hash TEXT, role TEXT)")
            legacy.execute("INSERT INTO users VALUES (1, 'admin', ?, 'admin')", (
                bcrypt.hashpw(b"admin", bcrypt.gensalt()).decode(),
            ))
            legacy.execute("""CREATE TABLE totp_recovery_codes (
                username TEXT NOT NULL REFERENCES users(username) ON DELETE CASCADE,
                code_hash TEXT NOT NULL, created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
                PRIMARY KEY(username, code_hash))""")
        legacy.close()

        db = SecureDatabase(path)
        self.addCleanup(db.close)
        self.assertEqual(db.conn.execute("SELECT COUNT(*) FROM users WHERE username='admin'").fetchone()[0], 0)
        self.assertEqual(db.conn.execute("PRAGMA foreign_key_list(totp_recovery_codes)").fetchall(), [])

    def test_vault_round_trip_and_text_parser(self):
        db = self.database("vault.db")
        self.addCleanup(db.close)
        source = os.path.join(self.tempdir.name, "source.txt")
        with open(source, "w", encoding="utf-8") as file:
            file.write("local confidential sample")
        self.assertEqual(DocumentParser.extract_text(source), ("local confidential sample", False))
        vault = SecureVault(db, os.path.join(self.tempdir.name, "vault"))
        self.assertTrue(vault.store_document(source, "Legal", "rootuser"))
        vault_id = db.conn.execute("SELECT id FROM document_vault").fetchone()[0]
        output_dir = os.path.join(self.tempdir.name, "out")
        os.makedirs(output_dir)
        self.assertTrue(vault.retrieve_document(vault_id, output_dir))
        with open(os.path.join(output_dir, "source.txt"), encoding="utf-8") as file:
            self.assertEqual(file.read(), "local confidential sample")

    def test_document_parsers_cover_pdf_docx_xlsx_csv_and_eml(self):
        import fitz
        import docx
        import openpyxl

        pdf_path = os.path.join(self.tempdir.name, "sample.pdf")
        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "PDF sample")
        pdf.save(pdf_path)
        pdf.close()
        self.assertIn("PDF sample", DocumentParser.extract_text(pdf_path)[0])

        docx_path = os.path.join(self.tempdir.name, "sample.docx")
        document = docx.Document()
        document.add_paragraph("DOCX sample")
        document.save(docx_path)
        self.assertIn("DOCX sample", DocumentParser.extract_text(docx_path)[0])

        xlsx_path = os.path.join(self.tempdir.name, "sample.xlsx")
        workbook = openpyxl.Workbook()
        workbook.active["A1"] = "XLSX sample"
        workbook.save(xlsx_path)
        self.assertIn("XLSX sample", DocumentParser.extract_text(xlsx_path)[0])

        csv_path = os.path.join(self.tempdir.name, "sample.csv")
        with open(csv_path, "w", encoding="utf-8") as file:
            file.write("name,value\nAda,CSV sample\n")
        self.assertIn("CSV sample", DocumentParser.extract_text(csv_path)[0])

        eml_path = os.path.join(self.tempdir.name, "sample.eml")
        with open(eml_path, "w", encoding="utf-8") as file:
            file.write("Subject: sample\n\nEML sample body")
        self.assertIn("EML sample body", DocumentParser.extract_text(eml_path)[0])

    def test_ocr_missing_dependencies_does_not_mark_image_clean(self):
        from PIL import Image
        import builtins

        image_path = os.path.join(self.tempdir.name, "ocr-fixture.png")
        Image.new("RGB", (64, 32), "white").save(image_path)
        real_import = builtins.__import__

        def without_cv2(name, *args, **kwargs):
            if name == "cv2":
                raise ImportError("test: OpenCV unavailable")
            return real_import(name, *args, **kwargs)

        with patch("builtins.__import__", side_effect=without_cv2):
            with self.assertLogs(level="ERROR"):
                with self.assertRaisesRegex(RuntimeError, "Image OCR is unavailable"):
                    DocumentParser.extract_text(image_path)

    def test_batch_scan_exports_redacted_text_and_pdf_report(self):
        from gui.batch_worker import BatchWorker

        input_dir = os.path.join(self.tempdir.name, "input")
        output_dir = os.path.join(self.tempdir.name, "output")
        os.makedirs(input_dir)
        os.makedirs(output_dir)
        with open(os.path.join(input_dir, "people.txt"), "w", encoding="utf-8") as file:
            file.write("Contact 0551234567")
        db = self.database("batch.db")
        self.addCleanup(db.close)
        worker = BatchWorker(
            input_dir,
            output_dir,
            AlgComplianceEngine(use_nlp=False),
            "full_mask",
            db=db,
        )
        worker.start()
        self.assertTrue(worker.wait(30000), "batch worker did not finish within 30 seconds")
        self.assertTrue(os.path.isfile(os.path.join(output_dir, "people.txt.anonymized.txt")))
        self.assertTrue(os.path.isfile(os.path.join(output_dir, "people.txt.report.pdf")))
        with open(os.path.join(output_dir, "people.txt.anonymized.txt"), encoding="utf-8-sig") as result:
            self.assertNotIn("0551234567", result.read())
        self.assertEqual(len(db.get_scan_history()), 1)

    def test_original_format_redactors_remove_detected_pii_and_verify(self):
        import csv
        import docx
        import fitz
        import openpyxl
        from engine.csv_redactor import redact_csv
        from engine.original_format_verifier import verify_original_format
        from engine.pdf_redactor import redact_pdf
        from engine.spreadsheet_redactor import redact_xlsx
        from engine.word_redactor import redact_docx

        engine = AlgComplianceEngine(use_nlp=False)
        phone = "0551234567"
        fixtures = []

        csv_in, csv_out = (os.path.join(self.tempdir.name, f"sample.{name}.csv") for name in ("in", "out"))
        with open(csv_in, "w", encoding="utf-8", newline="") as source:
            csv.writer(source).writerows([["Name", "Phone"], ["Synthetic", phone]])
        fixtures.append((".csv", redact_csv(csv_in, csv_out, engine, "full_mask"), csv_out))

        docx_in, docx_out = (os.path.join(self.tempdir.name, f"sample.{name}.docx") for name in ("in", "out"))
        document = docx.Document()
        document.add_paragraph(f"Contact {phone}")
        document.save(docx_in)
        fixtures.append((".docx", redact_docx(docx_in, docx_out, engine, "full_mask"), docx_out))

        xlsx_in, xlsx_out = (os.path.join(self.tempdir.name, f"sample.{name}.xlsx") for name in ("in", "out"))
        workbook = openpyxl.Workbook()
        workbook.active["A1"] = phone
        workbook.save(xlsx_in)
        workbook.close()
        fixtures.append((".xlsx", redact_xlsx(xlsx_in, xlsx_out, engine, "full_mask"), xlsx_out))

        pdf_in, pdf_out = (os.path.join(self.tempdir.name, f"sample.{name}.pdf") for name in ("in", "out"))
        document = fitz.open()
        page = document.new_page()
        page.insert_text((72, 72), f"Contact {phone}")
        document.save(pdf_in)
        document.close()
        fixtures.append((".pdf", redact_pdf(pdf_in, pdf_out, engine, "full_mask"), pdf_out))

        for extension, result, output_path in fixtures:
            with self.subTest(extension=extension):
                verification = verify_original_format(output_path, extension, result.scan_result, engine)
                self.assertTrue(verification["passed"], verification)

    def test_cron_schedule_uses_expression_and_last_run(self):
        scheduler = TaskScheduler.__new__(TaskScheduler)
        scheduler._is_running = True
        yesterday = datetime.now() - timedelta(days=1)
        task = (1, "daily", "", "0 2 * * *", "full_mask", 1, yesterday.strftime("%Y-%m-%d %H:%M:%S"), 0, 0, yesterday.strftime("%Y-%m-%d %H:%M:%S"))
        self.assertTrue(scheduler._should_run_task(task))
        invalid = task[:3] + ("not cron",) + task[4:]
        with self.assertLogs(level="ERROR"):
            self.assertFalse(scheduler._should_run_task(invalid))

    def test_scheduled_scan_logs_clean_files_and_detects_content_changes(self):
        path = os.path.join(self.tempdir.name, "scheduled")
        os.makedirs(path)
        source = os.path.join(path, "record.txt")
        with open(source, "w", encoding="utf-8") as file:
            file.write("No personal data here")
        db = self.database("schedule.db")
        self.addCleanup(db.close)
        db.save_scheduled_task("daily", path, "0 2 * * *", "legal_mask", "adminuser")
        task_id = db.conn.execute("SELECT id FROM scheduled_tasks").fetchone()[0]
        scheduler = TaskScheduler(db, AlgComplianceEngine(use_nlp=False))
        scheduler.db = db
        scheduler._execute_task(task_id, "daily", path, "legal_mask")
        self.assertEqual(len(db.get_scan_history()), 1)
        initial_mtime = os.stat(source).st_mtime
        scheduler._execute_task(task_id, "daily", path, "legal_mask")
        self.assertEqual(len(db.get_scan_history()), 1)
        with open(source, "w", encoding="utf-8") as file:
            file.write("Updated content, still no personal data")
        os.utime(source, (initial_mtime, initial_mtime))
        scheduler._execute_task(task_id, "daily", path, "legal_mask")
        self.assertEqual(len(db.get_scan_history()), 2)

    def test_admin_desktop_shell_constructs_and_closes(self):
        from PyQt6.QtWidgets import QApplication, QDialog
        from gui.login_dialog import LoginDialog
        from gui.main_window import AlgPIIMainWindow
        from gui.settings_dialog import SettingsDialog

        app = QApplication.instance() or QApplication([])
        db = self.database("ui.db")
        self.addCleanup(db.close)
        db.create_user("adminuser", "a-secure-password", "admin")
        settings_dialog = SettingsDialog(db, "admin")
        self.assertTrue(settings_dialog.retention_list)
        db.save_retention_policy("scan_history", 180)
        settings_dialog._refresh_retention_policies()
        self.assertEqual(settings_dialog.retention_list.count(), 1)
        settings_dialog.close()

        def accept_login(dialog):
            dialog.authenticated_role = "admin"
            dialog.authenticated_username = "adminuser"
            dialog.authenticated_department = None
            return QDialog.DialogCode.Accepted

        with patch.object(LicenseManager, "is_activated", return_value=True), \
             patch.object(LoginDialog, "exec", accept_login):
            window = AlgPIIMainWindow(db)
        try:
            self.assertEqual(window.content_stack.count(), 15)
            self.assertEqual(window.user_role, "admin")
            scan_result = window.engine.anonymize("نص تجريبي بلا بيانات شخصية", "full_mask")
            window.on_scan_complete(scan_result)
            self.assertGreaterEqual(window.badges_layout.count(), 1)
            self.assertEqual(db.get_scan_history()[0][3], 0)
            with patch("gui.pia_page.PIAWizard.exec", return_value=QDialog.DialogCode.Rejected):
                window.pia_page.on_new_pia()
            db.create_incident("اختبار", "حادث تجريبي", "MEDIUM", "test", 1, "adminuser")
            with self.assertNoLogs(level="ERROR"):
                window.dashboard_page.refresh_data()
            for page_index in range(window.content_stack.count()):
                window._switch_page(page_index)
                app.processEvents()
            self.assertTrue(hasattr(window, "calendar_page"))
            self.assertTrue(hasattr(window, "training_page"))
            self.assertTrue(hasattr(window, "report_template_page"))
        finally:
            window.scheduler.stop()
            window.scheduler.wait(10000)
            window.tray.hide()
            db.close()
            app.processEvents()

    def test_status_indicator_pulse_opacity_stays_in_color_range(self):
        from PyQt6.QtWidgets import QApplication
        from gui.widgets import StatusIndicator

        app = QApplication.instance() or QApplication([])
        indicator = StatusIndicator()
        indicator.set_state("loading")
        for _ in range(100):
            indicator._animate()
            self.assertGreaterEqual(indicator.opacity, 0.0)
            self.assertLessEqual(indicator.opacity, 1.0)
        indicator.timer.stop()

    def test_business_workflows_pages_and_export_artifacts(self):
        from PyQt6.QtWidgets import QApplication, QMessageBox
        from engine.cross_linker import CrossLinker
        from engine.dsar_manager import DSARManager
        from engine.evidence_packager import EvidencePackager
        from engine.exporters import ReportExporter
        from engine.pdf_exporter import PDFExporter
        from gui.consent_page import ConsentPageWidget
        from gui.data_flow_page import DataFlowPageWidget
        from gui.dsar_page import DSARPageWidget
        from gui.incidents_page import IncidentsPageWidget
        from gui.pia_page import PIAPageWidget
        from gui.pia_wizard import PIAWizard
        from gui.report_template_page import ReportTemplatePageWidget
        from gui.training_page import TrainingPageWidget

        app = QApplication.instance() or QApplication([])
        db = self.database("workflows.db")
        self.addCleanup(db.close)
        db.save_department("HR")
        db.save_department("Legal")
        hr_id, legal_id = [row[0] for row in db.get_departments()]

        engine = AlgComplianceEngine(use_nlp=False)
        scanned = engine.anonymize("Contact demo.person@example.test or 0551234567", "full_mask")
        scan_id = db.log_scan("synthetic.txt", len(scanned.entities), "HIGH", "full_mask", {e.entity_type: 1 for e in scanned.entities}, department="HR")
        CrossLinker(db).index_scan_result(scan_id, "synthetic.txt", scanned.entities, "HR")
        report = engine.generate_report(scanned.scan_result)

        # All report formats are created from the same synthetic scan.
        paths = {
            "pdf": os.path.join(self.tempdir.name, "report.pdf"),
            "docx": os.path.join(self.tempdir.name, "report.docx"),
            "xlsx": os.path.join(self.tempdir.name, "report.xlsx"),
            "json": os.path.join(self.tempdir.name, "report.json"),
            "xml": os.path.join(self.tempdir.name, "report.xml"),
            "csv": os.path.join(self.tempdir.name, "report.csv"),
        }
        PDFExporter.export_report(scanned, report, paths["pdf"])
        ReportExporter.save_docx(scanned.scan_result, report, paths["docx"])
        ReportExporter.save_xlsx(scanned.scan_result, paths["xlsx"])
        with open(paths["json"], "w", encoding="utf-8") as file:
            file.write(ReportExporter.to_json(scanned.scan_result, report))
        with open(paths["xml"], "w", encoding="utf-8") as file:
            file.write(ReportExporter.to_xml(scanned.scan_result, report))
        with open(paths["csv"], "w", encoding="utf-8") as file:
            file.write(ReportExporter.to_csv(scanned.scan_result))
        self.assertTrue(all(os.path.getsize(path) > 0 for path in paths.values()))

        # DSAR search and response export.
        email = next(entity.text for entity in scanned.entities if entity.entity_type == "EMAIL")
        request_id = db.create_dsar_request("Synthetic Requester", "requester@example.test", email)
        self.assertEqual(DSARManager(db).find_subject_data(email)[0]["document_name"], "synthetic.txt")
        dsar_path = os.path.join(self.tempdir.name, "dsar.pdf")
        DSARManager(db).generate_dsar_report(request_id, email, dsar_path)
        db.update_dsar_status(request_id, "CLOSED", "Synthetic test complete")
        self.assertTrue(os.path.getsize(dsar_path) > 0)
        self.assertEqual(db.get_dsar_requests()[0]["status"], "CLOSED")

        # PIA form persistence and PDF report.
        wizard = PIAWizard(db, "HR", "adminuser")
        wizard.project_name_edit.setText("Synthetic PIA")
        wizard.purpose_edit.setPlainText("Test only")
        wizard.risk_combo.setCurrentIndex(2)
        wizard.mitigation_edit.setPlainText("Encrypt and restrict access")
        with patch.object(QMessageBox, "information"):
            wizard.accept()
        pia_page = PIAPageWidget(db, None, "adminuser")
        self.assertEqual(pia_page.table.rowCount(), 1)
        pia_row = db.conn.execute("SELECT project_name, department, assessor_name, data_types_json, processing_purpose, risk_level, mitigation_steps, created_at FROM pia_assessments").fetchone()
        pia_pdf = os.path.join(self.tempdir.name, "pia.pdf")
        self.assertTrue(PDFExporter.export_pia(dict(zip(("project_name", "department", "assessor_name", "data_types", "processing_purpose", "risk_level", "mitigation_steps", "created_at"), pia_row)), pia_pdf))
        self.assertTrue(os.path.getsize(pia_pdf) > 0)

        # Consent, transfers, data-flow stats, incidents and evidence package.
        db.conn.execute("INSERT INTO consent_registry (subject_id, subject_name, consent_type, status, source) VALUES (?, ?, ?, ?, ?)", (email, "Synthetic Person", "DATA_PROCESSING", "OPT_IN", "WEB_FORM"))
        db.conn.commit()
        consent_page = ConsentPageWidget(db)
        self.assertEqual(consent_page.table.rowCount(), 1)
        transfer_id = db.create_transfer(hr_id, legal_id, scan_id, scan_id, "full_mask", "Synthetic test transfer", "adminuser", "legal")
        db.update_transfer_status(transfer_id, "APPROVED")
        flow_page = DataFlowPageWidget(db)
        flow_page.load_data()
        self.assertEqual(flow_page.transfers_table.rowCount(), 1)
        incident_id = db.create_incident("Synthetic incident", "Test evidence export", "MEDIUM", "scan_alert", scan_id, "adminuser")
        incidents_page = IncidentsPageWidget(db)
        self.assertEqual(incidents_page.columns["OPEN"].content_layout.count(), 1)
        evidence_path = EvidencePackager(db).generate_package(incident_id, self.tempdir.name)
        self.assertTrue(os.path.isfile(evidence_path))

        # Training result and report-template list refresh.
        quizzes = db.get_training_quizzes()
        self.assertGreater(len(quizzes), 0)
        db.save_training_result("adminuser", str(quizzes[0][0]), 100, True)
        training_page = TrainingPageWidget(db, "adminuser")
        self.assertGreaterEqual(training_page.quizzes_table.rowCount(), 1)
        self.assertEqual(training_page.results_table.rowCount(), 1)
        db.save_report_template("Synthetic report", "Test template", "[]", "", True, "adminuser")
        templates_page = ReportTemplatePageWidget(db)
        self.assertEqual(templates_page.templates_list.count(), 1)

        # Pages exercise their data loaders against populated synthetic records.
        dsar_page = DSARPageWidget(db)
        dsar_page.load_data()
        self.assertEqual(dsar_page.table.rowCount(), 1)
        for widget in (wizard, pia_page, consent_page, flow_page, incidents_page, training_page, templates_page, dsar_page):
            widget.close()
        app.processEvents()


if __name__ == "__main__":
    unittest.main()
