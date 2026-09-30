import os
import csv
import logging
import tempfile
import shutil

from typing import Tuple

class DocumentParser:
    """
    Enterprise Document Parser for Alg-PII Engine.
    Extracts plain text from various file formats (TXT, PDF, DOCX, XLSX, CSV).
    Returns (extracted_text, ocr_used)
    """
    _ocr_enabled = True
    _ocr_language = "ara+eng"
    _tesseract_cmd = ""

    @classmethod
    def configure_ocr(cls, enabled=True, language="ara+eng", tesseract_cmd=""):
        cls._ocr_enabled = bool(enabled)
        cls._ocr_language = language or "ara+eng"
        cls._tesseract_cmd = (tesseract_cmd or "").strip()

    @classmethod
    def get_ocr_config(cls):
        return cls._ocr_enabled, cls._ocr_language, cls._tesseract_cmd

    @classmethod
    def validate_ocr_configuration(cls, language=None, tesseract_cmd=None):
        try:
            import pytesseract
            configured_path = (tesseract_cmd if tesseract_cmd is not None else cls._tesseract_cmd) or ""
            if configured_path:
                if not os.path.isfile(configured_path):
                    return False, "مسار tesseract.exe غير موجود."
                pytesseract.pytesseract.tesseract_cmd = configured_path
            else:
                path_from_env = shutil.which("tesseract")
                if not path_from_env:
                    return False, "لم يُعثر على Tesseract في PATH؛ حدد مسار tesseract.exe."
                pytesseract.pytesseract.tesseract_cmd = path_from_env

            pytesseract.get_tesseract_version()
            requested = (language if language is not None else cls._ocr_language) or "ara+eng"
            available = set(pytesseract.get_languages(config=""))
            missing = [item for item in requested.split("+") if item not in available]
            if missing:
                return False, "حزم اللغة غير مثبتة: " + ", ".join(missing)
            return True, "Tesseract وحزم اللغة المطلوبة جاهزة."
        except Exception as exc:
            return False, f"تعذر تشغيل Tesseract: {exc}"

    @staticmethod
    def extract_text(file_path: str) -> Tuple[str, bool]:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        ext = os.path.splitext(file_path)[1].lower()

        try:
            if ext == '.txt':
                return DocumentParser._parse_txt(file_path), False
            elif ext == '.pdf':
                return DocumentParser._parse_pdf(file_path)
            elif ext == '.docx':
                return DocumentParser._parse_docx(file_path), False
            elif ext == '.xlsx':
                return DocumentParser._parse_xlsx(file_path), False
            elif ext in ['.png', '.jpg', '.jpeg', '.tiff', '.bmp']:
                return DocumentParser._parse_image(file_path), True
            elif ext == '.csv':
                return DocumentParser._parse_csv(file_path), False
            elif ext == '.eml':
                return DocumentParser._parse_eml(file_path), False
            elif ext == '.msg':
                return DocumentParser._parse_msg(file_path), False
            else:
                raise ValueError(f"Unsupported file format: {ext}")
        except Exception as e:
            logging.error(f"Error parsing {file_path}: {e}")
            raise RuntimeError(f"Failed to extract text from {file_path}: {str(e)}")

    @staticmethod
    def _parse_eml(file_path: str) -> str:
        import email
        from email import policy
        
        text_content = []
        with open(file_path, 'rb') as f:
            msg = email.message_from_binary_file(f, policy=policy.default)
            
        text_content.append(f"Subject: {msg.get('subject', '')}")
        text_content.append(f"From: {msg.get('from', '')}")
        text_content.append(f"To: {msg.get('to', '')}")
        text_content.append(f"Date: {msg.get('date', '')}")
        text_content.append("-" * 40)
        
        for part in msg.walk():
            if part.get_content_type() == 'text/plain':
                try:
                    body = part.get_content()
                    text_content.append(body)
                except Exception:
                    pass
                    
        return "\n".join(text_content)

    @staticmethod
    def _parse_msg(file_path: str) -> str:
        try:
            import extract_msg
            msg = extract_msg.Message(file_path)
            
            text_content = []
            text_content.append(f"Subject: {msg.subject}")
            text_content.append(f"From: {msg.sender}")
            text_content.append(f"To: {msg.to}")
            text_content.append(f"Date: {msg.date}")
            text_content.append("-" * 40)
            
            if msg.body:
                text_content.append(msg.body)
                
            return "\n".join(text_content)
        except ImportError:
            logging.warning("extract-msg not installed. Cannot parse .msg files.")
            raise RuntimeError("Install the extract-msg package to read Outlook .msg files.")

    @staticmethod
    def _parse_txt(file_path: str) -> str:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()

    @staticmethod
    def _parse_pdf(file_path: str) -> Tuple[str, bool]:
        import fitz  # PyMuPDF
        text_content = []
        is_scanned = True
        ocr_used = False
        
        with fitz.open(file_path) as doc:
            for page in doc:
                text = page.get_text().strip()
                if text:
                    is_scanned = False
                    text_content.append(text)
                elif DocumentParser._ocr_enabled:
                    # Try OCR on the page if it's empty (scanned image)
                    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
                    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as temp_image:
                        img_path = temp_image.name
                    pix.save(img_path)
                    try:
                        ocr_text = DocumentParser._parse_image(img_path)
                        if ocr_text:
                            text_content.append(ocr_text)
                            ocr_used = True
                    finally:
                        if os.path.exists(img_path):
                            os.remove(img_path)
                elif page.get_images(full=True):
                    raise RuntimeError("هذا PDF ممسوح ضوئيًا وOCR معطل. فعّل OCR من الإعدادات لفحصه.")
                            
        return "\n\n".join(text_content), ocr_used

    @staticmethod
    def _parse_image(file_path: str) -> str:
        if not DocumentParser._ocr_enabled:
            raise RuntimeError("تم تعطيل OCR من الإعدادات؛ لا يمكن استخراج النص من الصورة.")
        try:
            import cv2
            import pytesseract
            from PIL import Image
        except ImportError as exc:
            logging.error("OCR dependencies are missing; refusing to treat the image as text-free.")
            raise RuntimeError(
                "Image OCR is unavailable. Install opencv-python-headless and pytesseract, "
                "plus the Tesseract executable with Arabic and English language data."
            ) from exc

        # Preprocess with OpenCV before OCR. Any missing executable/language pack
        # error is propagated so a scan cannot silently report an image as clean.
        img = cv2.imread(file_path)
        if img is None:
            raise RuntimeError("OCR could not open this image file.")

        gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
        denoised = cv2.fastNlMeansDenoising(gray, h=30)
        _, thresh = cv2.threshold(denoised, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
        pil_img = Image.fromarray(thresh)
        configured_path = DocumentParser._tesseract_cmd
        if configured_path:
            if not os.path.isfile(configured_path):
                raise RuntimeError("مسار Tesseract المحدد غير موجود: " + configured_path)
            pytesseract.pytesseract.tesseract_cmd = configured_path
        else:
            path_from_env = shutil.which("tesseract")
            if not path_from_env:
                raise RuntimeError("Tesseract OCR غير موجود في PATH. حدّد مسار tesseract.exe من الإعدادات.")
            pytesseract.pytesseract.tesseract_cmd = path_from_env
        return pytesseract.image_to_string(pil_img, lang=DocumentParser._ocr_language).strip()

    @staticmethod
    def _parse_docx(file_path: str) -> str:
        import docx
        doc = docx.Document(file_path)
        return "\n".join([paragraph.text for paragraph in doc.paragraphs])

    @staticmethod
    def _parse_xlsx(file_path: str) -> str:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, data_only=True)
        text_content = []
        for sheetname in wb.sheetnames:
            sheet = wb[sheetname]
            for row in sheet.iter_rows(values_only=True):
                # Filter out None values and convert everything to string
                row_texts = [str(cell) for cell in row if cell is not None]
                if row_texts:
                    text_content.append(" \t ".join(row_texts))
        return "\n".join(text_content)

    @staticmethod
    def _parse_csv(file_path: str) -> str:
        text_content = []
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            reader = csv.reader(f)
            for row in reader:
                if row:
                    text_content.append(" \t ".join(row))
        return "\n".join(text_content)
