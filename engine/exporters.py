import csv
import io
import json
import xml.etree.ElementTree as ET
from collections import Counter

from .regex_detector import ComplianceReport, ScanResult


def _entity_counts(result: ScanResult) -> dict:
    return dict(Counter(entity.entity_type for entity in result.entities))


def _csv_safe(value) -> str:
    """Prevent spreadsheet programs from evaluating exported text as a formula."""
    value = str(value or "")
    if value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _safe_report_data(result: ScanResult, report: ComplianceReport, anonymized_text: str) -> dict:
    """Build a shareable report without exporting the original entity values."""
    return {
        "risk_level": report.risk_level,
        "sensitivity_label": report.sensitivity_label,
        "entity_counts": _entity_counts(result),
        "recommendations": list(report.recommendations),
        "clusters": report.clusters,
        "processed_text": anonymized_text,
    }


class ReportExporter:
    """Export scan summaries and anonymized text without raw detected values."""

    @staticmethod
    def to_json(result: ScanResult, report: ComplianceReport, anonymized_text: str = "") -> str:
        return json.dumps(
            _safe_report_data(result, report, anonymized_text),
            ensure_ascii=False,
            indent=2,
        )

    @staticmethod
    def to_xml(result: ScanResult, report: ComplianceReport, anonymized_text: str = "") -> str:
        data = _safe_report_data(result, report, anonymized_text)
        root = ET.Element("ComplianceReport")
        ET.SubElement(root, "RiskLevel").text = data["risk_level"]
        ET.SubElement(root, "SensitivityLabel").text = data["sensitivity_label"]

        counts = ET.SubElement(root, "EntityCounts")
        for entity_type, count in data["entity_counts"].items():
            item = ET.SubElement(counts, "EntityType", name=entity_type)
            item.text = str(count)

        recommendations = ET.SubElement(root, "Recommendations")
        for recommendation in data["recommendations"]:
            ET.SubElement(recommendations, "Recommendation").text = recommendation

        ET.SubElement(root, "ProcessedText").text = data["processed_text"]
        xml_str = ET.tostring(root, encoding="utf-8")
        try:
            import xml.dom.minidom
            return xml.dom.minidom.parseString(xml_str).toprettyxml(indent="  ")
        except Exception:
            return xml_str.decode("utf-8")

    @staticmethod
    def to_csv(result: ScanResult, report: ComplianceReport = None, anonymized_text: str = "") -> str:
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(["Record", "Value"])
        if report is not None:
            writer.writerow(["Risk level", _csv_safe(report.risk_level)])
            writer.writerow(["Sensitivity label", _csv_safe(report.sensitivity_label)])
        for entity_type, count in _entity_counts(result).items():
            writer.writerow([f"Detected count: {entity_type}", count])
        writer.writerow(["Processed text", _csv_safe(anonymized_text)])
        return output.getvalue()

    @staticmethod
    def save_docx(result: ScanResult, report: ComplianceReport, file_path: str, anonymized_text: str = ""):
        try:
            from docx import Document
        except ImportError as exc:
            raise ImportError("python-docx is not installed. Please 'pip install python-docx'.") from exc

        document = Document()
        document.add_heading("تقرير امتثال Alg-PII Engine", 0)
        document.add_heading("الملخص", level=1)
        document.add_paragraph(f"مستوى الخطر: {report.risk_level}")
        document.add_paragraph(f"التصنيف: {report.sensitivity_label}")

        document.add_heading("أعداد أنواع البيانات المكتشفة", level=1)
        counts = _entity_counts(result)
        if counts:
            table = document.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "نوع البيانات"
            table.rows[0].cells[1].text = "العدد"
            for entity_type, count in counts.items():
                cells = table.add_row().cells
                cells[0].text = entity_type
                cells[1].text = str(count)
        else:
            document.add_paragraph("لم تُكتشف كيانات وفق قواعد الفحص المفعلة.")

        document.add_heading("التوصيات", level=1)
        for recommendation in report.recommendations:
            document.add_paragraph(recommendation, style="List Bullet")

        document.add_heading("النص المعالج", level=1)
        document.add_paragraph(anonymized_text)
        document.save(file_path)

    @staticmethod
    def save_xlsx(result: ScanResult, file_path: str, report: ComplianceReport = None, anonymized_text: str = ""):
        try:
            from openpyxl import Workbook
        except ImportError as exc:
            raise ImportError("openpyxl is not installed. Please 'pip install openpyxl'.") from exc

        workbook = Workbook()
        summary = workbook.active
        summary.title = "Summary"
        summary.append(["Metric", "Value"])
        if report is not None:
            summary.append(["Risk level", report.risk_level])
            summary.append(["Sensitivity label", report.sensitivity_label])

        entities = workbook.create_sheet("Entity Counts")
        entities.append(["Type", "Count"])
        for entity_type, count in _entity_counts(result).items():
            entities.append([entity_type, count])

        processed = workbook.create_sheet("Processed Text")
        for line in anonymized_text.splitlines() or [""]:
            cell = processed.cell(row=processed.max_row + 1, column=1, value=line)
            # Keep untrusted content as text, even when it begins with '='.
            cell.data_type = "s"

        workbook.save(file_path)
